# RefCheck CGB/UNESP - Verificador de Referências

import requests
import streamlit as st
import re
import fitz  # PyMuPDF para PDF
import docx
from io import StringIO, BytesIO
from fpdf import FPDF

st.set_page_config(page_title="RefCheck - CGB UNESP", layout="centered")

# Funções de verificação

def buscar_por_doi(doi):
    url = f"https://api.crossref.org/works/{doi}"
    resposta = requests.get(url)
    if resposta.status_code == 200:
        dados = resposta.json()
        titulo = dados['message'].get('title', [''])[0]
        autores = [autor.get('family', '') for autor in dados['message'].get('author', [])]
        return {
            'status': 'confirmado',
            'titulo': titulo,
            'autores': autores,
            'link': f"https://doi.org/{doi}"
        }
    else:
        return {
            'status': 'nao_encontrado',
            'mensagem': f"DOI {doi} não encontrado na CrossRef."
        }

def buscar_por_isbn(isbn):
    url = f"https://openlibrary.org/api/books?bibkeys=ISBN:{isbn}&format=json&jscmd=data"
    resposta = requests.get(url)
    dados = resposta.json()
    if f"ISBN:{isbn}" in dados:
        livro = dados[f"ISBN:{isbn}"]
        titulo = livro.get('title', '')
        autores = [autor.get('name') for autor in livro.get('authors', [])]
        return {
            'status': 'confirmado',
            'titulo': titulo,
            'autores': autores,
            'link': livro.get('url')
        }
    else:
        return {
            'status': 'nao_encontrado',
            'mensagem': f"ISBN {isbn} não encontrado no OpenLibrary."
        }

# Funções auxiliares

def extrair_dois(texto):
    padrao = re.compile(r"10\.\d{4,9}/[-._;()/:A-Z0-9]+", re.IGNORECASE)
    return padrao.findall(texto)

def extrair_isbns(texto):
    padrao = re.compile(r"\b(?:ISBN[- ]*)?(?=\d{10,13}\b)[0-9\-]{10,17}\b")
    return [re.sub(r'[^0-9X]', '', isbn) for isbn in padrao.findall(texto)]

def ler_pdf(arquivo):
    texto = ""
    with fitz.open(stream=arquivo.read(), filetype="pdf") as doc:
        for pagina in doc:
            texto += pagina.get_text()
    return texto

def ler_docx(arquivo):
    doc = docx.Document(arquivo)
    return "\n".join([par.text for par in doc.paragraphs])

def gerar_relatorio_pdf(resultados):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Relatório de Verificação de Referências - CGB UNESP", ln=True, align='C')
    pdf.ln(10)
    for item in resultados:
        pdf.set_font("Arial", 'B', size=11)
        pdf.cell(200, 10, txt=f"{item['tipo'].upper()} encontrado: {item['entrada']}", ln=True)
        if item['status'] == 'confirmado':
            pdf.set_font("Arial", size=10)
            pdf.multi_cell(0, 10, txt=f"Título: {item['titulo']}\nLink: {item['link']}\n")
        else:
            pdf.set_font("Arial", size=10)
            pdf.cell(200, 10, txt="Referência não encontrada.", ln=True)
        pdf.ln(5)
    buffer = BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer

# Interface Streamlit
st.title("🔎 RefCheck - Verificador de Referências | CGB UNESP")
st.markdown("Verifique automaticamente se as referências bibliográficas citadas em seu texto existem e estão corretas.")

entrada = st.radio("Como você quer inserir as referências?", ["Colar texto", "Enviar arquivo"])

texto = ""
if entrada == "Colar texto":
    texto = st.text_area("Cole aqui o conteúdo do seu artigo ou lista de referências")
elif entrada == "Enviar arquivo":
    arquivo = st.file_uploader("Envie um arquivo (.txt, .pdf, .docx)", type=["txt", "pdf", "docx"])
    if arquivo:
        if arquivo.name.endswith(".txt"):
            texto = arquivo.read().decode("utf-8", errors='ignore')
        elif arquivo.name.endswith(".pdf"):
            texto = ler_pdf(arquivo)
        elif arquivo.name.endswith(".docx"):
            texto = ler_docx(arquivo)

if texto:
    st.markdown("### 🔎 Verificando referências...")
    dois = extrair_dois(texto)
    isbns = extrair_isbns(texto)
    resultados = []

    for doi in dois:
        resultado = buscar_por_doi(doi)
        resultado.update({"entrada": doi, "tipo": "doi"})
        resultados.append(resultado)

    for isbn in isbns:
        resultado = buscar_por_isbn(isbn)
        resultado.update({"entrada": isbn, "tipo": "isbn"})
        resultados.append(resultado)

    for r in resultados:
        if r['status'] == 'confirmado':
            st.success(f"{r['tipo'].upper()} confirmado: {r['entrada']}")
            st.markdown(f"- **Título:** {r['titulo']}")
            st.markdown(f"- **Link:** {r['link']}")
        else:
            st.warning(f"{r['tipo'].upper()} não encontrado: {r['entrada']}")

    if resultados:
        buffer = gerar_relatorio_pdf(resultados)
        st.download_button("📄 Baixar Relatório em PDF", data=buffer, file_name="relatorio_referencias.pdf", mime="application/pdf")
    else:
        st.info("Nenhum DOI ou ISBN encontrado no texto enviado.")
